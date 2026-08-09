from __future__ import annotations

from pathlib import Path
from typing import cast

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from spec_pipeline.exporters.customer_fields import (
    customer_description,
    customer_note,
    normalize_customer_unit,
)
from spec_pipeline.exporters.excel_safe import excel_safe_text, excel_safe_value
from spec_pipeline.io import load_rows_from_jsonl
from spec_pipeline.models import SampleDefinition, SpecificationRow
from spec_pipeline.samples import load_ground_truth_rows
from spec_pipeline.workflows.normalize import default_output_path as default_normalized_output_path

HEADER = [
    "№ п/п",
    "Наименование",
    "Ед. изм.",
    "Кол-во",
    "Примечание",
]

SHEET_NAME = "Data"
CROSSED_OUT_STATUS = "crossed_out"
CROSSED_OUT_FONT_COLOR = "FF0000"
QUANTITY_NUMBER_FORMAT = "0.00"

DEFAULT_COLUMN_WIDTHS = {
    "A": 8,
    "B": 110,
    "C": 12,
    "D": 12,
    "E": 28,
}


def default_xlsx_output_path(project_root: Path, sample_id: str) -> Path:
    return project_root / "artifacts" / "xlsx" / f"{sample_id}.xlsx"


def row_to_excel_values(row: SpecificationRow) -> list[object | None]:
    return [
        excel_safe_text(row.pos),
        excel_safe_text(customer_description(row)),
        excel_safe_text(normalize_customer_unit(row.unit)),
        excel_safe_value(float(row.quantity) if row.quantity is not None else row.quantity_raw),
        excel_safe_text(customer_note(row)),
    ]


def write_rows_to_xlsx(rows: list[SpecificationRow], output_path: Path) -> Path:
    workbook = Workbook()
    worksheet = cast(Worksheet, workbook.active)
    worksheet.title = SHEET_NAME

    header_font = Font(bold=True)
    header_fill = PatternFill(fill_type="solid", fgColor="D9E2F3")
    grouping_fill = PatternFill(fill_type="solid", fgColor="F2F2F2")
    top_alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.append(HEADER)
    for cell in worksheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = top_alignment

    current_section: str | None = None
    current_subsection: str | None = None
    crossed_out_input_indexes = infer_crossed_out_input_indexes(rows)
    for input_index, row in enumerate(rows):
        if row.section and row.section != current_section:
            append_grouping_row(worksheet, row.section, grouping_fill, top_alignment, bold=True)
            current_section = row.section
            current_subsection = None
        if row.subsection and row.subsection != current_subsection:
            append_grouping_row(
                worksheet, row.subsection, grouping_fill, top_alignment, italic=True
            )
            current_subsection = row.subsection

        if row.is_grouping_row():
            continue

        worksheet.append(row_to_excel_values(row))
        current_row = worksheet.max_row
        for cell in worksheet[current_row]:
            cell.alignment = top_alignment
        if input_index in crossed_out_input_indexes:
            apply_crossed_out_font(worksheet, current_row)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:E{worksheet.max_row}"

    for column, width in DEFAULT_COLUMN_WIDTHS.items():
        worksheet.column_dimensions[column].width = width

    for column in ("D",):
        for cell in worksheet[column][1:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = QUANTITY_NUMBER_FORMAT

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def append_grouping_row(
    worksheet: Worksheet,
    label: str,
    fill: PatternFill,
    alignment: Alignment,
    *,
    bold: bool = False,
    italic: bool = False,
) -> None:
    worksheet.append([None, excel_safe_text(label), None, None, None])
    current_row = worksheet.max_row
    for cell in worksheet[current_row]:
        cell.fill = fill
        cell.alignment = alignment
    worksheet.cell(current_row, 2).font = Font(bold=bold, italic=italic)


def infer_crossed_out_input_indexes(rows: list[SpecificationRow]) -> set[int]:
    crossed: set[int] = set()
    previous_data_index: int | None = None
    for index, row in enumerate(rows):
        if row.status == CROSSED_OUT_STATUS:
            crossed.add(index)
            if is_text_continuation(row) and previous_data_index is not None:
                crossed.add(previous_data_index)
        if not row.is_grouping_row() and row.name is not None:
            previous_data_index = index
    return crossed


def is_text_continuation(row: SpecificationRow) -> bool:
    return row.name is not None and row.pos is None and row.unit is None and row.quantity is None


def apply_crossed_out_font(worksheet: Worksheet, row_number: int) -> None:
    for cell in worksheet[row_number]:
        source = cell.font
        cell.font = Font(
            name=source.name,
            sz=source.sz,
            b=source.b,
            i=source.i,
            vertAlign=source.vertAlign,
            underline=source.underline,
            strike=source.strike,
            color=CROSSED_OUT_FONT_COLOR,
            family=source.family,
            scheme=source.scheme,
        )


def export_sample_to_xlsx(
    *,
    sample: SampleDefinition,
    project_root: Path,
    output_path: str | Path | None = None,
    normalized_input_path: str | Path | None = None,
) -> tuple[Path, int]:
    if normalized_input_path is not None:
        rows = load_rows_from_jsonl(normalized_input_path)
    else:
        default_jsonl = default_normalized_output_path(project_root, sample.sample_id)
        rows = (
            load_rows_from_jsonl(default_jsonl)
            if default_jsonl.exists()
            else load_ground_truth_rows(sample)
        )

    resolved_output = (
        Path(output_path).expanduser()
        if output_path is not None
        else default_xlsx_output_path(project_root, sample.sample_id)
    )

    if output_path is not None and not resolved_output.parent.exists():
        raise FileNotFoundError(f"Output directory does not exist: {resolved_output.parent}")

    write_rows_to_xlsx(rows, resolved_output)
    return resolved_output, len(rows)


def export_jsonl_to_xlsx(
    *, input_jsonl_path: str | Path, output_path: str | Path
) -> tuple[Path, int]:
    rows = load_rows_from_jsonl(input_jsonl_path)
    resolved_output = Path(output_path).expanduser()
    if not resolved_output.parent.exists():
        raise FileNotFoundError(f"Output directory does not exist: {resolved_output.parent}")
    write_rows_to_xlsx(rows, resolved_output)
    return resolved_output, len(rows)
