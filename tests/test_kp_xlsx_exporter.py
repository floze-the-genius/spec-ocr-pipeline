from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path
from typing import cast

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from spec_pipeline.exporters.kp_xlsx import (
    MAIN_SHEET_NAME,
    build_kp_output_rows,
    commercial_unit_quantity,
    export_jsonl_to_kp_xlsx,
    kp_display_name,
    linear_quantity_to_surface_area,
)
from spec_pipeline.models import SpecificationRow


def _write_kp_template(tmp_path: Path) -> Path:
    workbook = Workbook()
    worksheet = cast(Worksheet, workbook.active)
    worksheet.title = MAIN_SHEET_NAME
    workbook.create_sheet("Лист2")
    workbook.create_sheet("Лист3")

    worksheet.column_dimensions["B"].width = 80
    worksheet.column_dimensions["K"].hidden = True
    worksheet.merge_cells("A7:A9")
    worksheet.merge_cells("B7:B9")
    worksheet.merge_cells("C7:C9")
    worksheet.merge_cells("D7:D9")
    worksheet.merge_cells("E7:F7")
    worksheet.merge_cells("E8:E9")
    worksheet.merge_cells("F8:F9")
    worksheet.merge_cells("G7:H7")
    worksheet.merge_cells("G8:G9")
    worksheet.merge_cells("H8:H9")
    worksheet.merge_cells("I7:I9")
    worksheet["A7"] = "№ п/п"
    worksheet["B7"] = "Наименование"
    worksheet["C7"] = "Ед. изм."
    worksheet["D7"] = "Кол-во"
    worksheet["E7"] = "Цена за единицу"
    worksheet["E8"] = "Материал"
    worksheet["F8"] = "Работа"
    worksheet["G7"] = "Сумма"
    worksheet["G8"] = "Материал"
    worksheet["H8"] = "Работа"
    worksheet["I7"] = "Примечание"
    for column in range(1, 10):
        cell = cast(Cell, worksheet.cell(10, column))
        cell.value = column
        cell.font = Font(bold=True)

    worksheet["B11"] = "Система вентиляции"
    worksheet["B11"].font = Font(bold=True)
    worksheet["B12"] = "ПВ7"
    worksheet["B12"].font = Font(bold=True)
    worksheet["A13"] = 1
    worksheet["B13"] = "Шаблонная строка"
    worksheet["C13"] = "шт"
    worksheet["D13"] = 1
    worksheet["G13"] = "=E13*D13"
    worksheet["H13"] = "=F13*D13"

    for row in range(20, 28):
        worksheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=7)
    worksheet["C20"] = "Итого стоимость оборудования и материалов:"
    worksheet["C21"] = "Итого стоимость оборудования:"
    worksheet["C22"] = "Итого стоимость материалов:"
    worksheet["C23"] = "Стоимость монтажных работ:"
    worksheet["C24"] = "Расходные и крепежные материалы:"
    worksheet["C25"] = "Транспортные и накладные расходы:"
    worksheet["C26"] = "ИТОГО:"
    worksheet["C27"] = "в т.ч. НДС 22%"
    worksheet["H20"] = "=SUM(G11:G19)"
    worksheet["H23"] = "=SUM(H11:H19)"
    worksheet.print_area = "A1:I30"

    path = tmp_path / "kp_template.xlsx"
    workbook.save(path)
    return path


def _write_kp_header_template(tmp_path: Path) -> Path:
    workbook = Workbook()
    worksheet = cast(Worksheet, workbook.active)
    worksheet.title = "Таблица"
    worksheet.merge_cells("A2:C2")
    worksheet.merge_cells("H2:J2")
    worksheet.merge_cells("A3:J3")
    worksheet.merge_cells("A4:J4")
    worksheet["A2"] = "согласовано                          м.п."
    worksheet["H2"] = "утверждаю                        м.п."
    worksheet["A3"] = 'Коммерческое предложение ООО "РОГА КОПЫТА" по объекту:'
    worksheet["A4"] = "ОБЪЕКТ МЖД"
    worksheet.row_dimensions[4].height = 40
    path = tmp_path / "9_КП_по_спец_№_8.xlsx"
    workbook.save(path)
    return path


def _write_rows_jsonl(tmp_path: Path) -> Path:
    rows = [
        {
            "sample_id": "sample",
            "page_ref": "1",
            "section": "ПВ7",
            "subsection": "КИПиА",
            "pos": "4",
            "name": "Датчик",
            "type_mark": "ТС-1",
            "code": None,
            "manufacturer": "СВОК",
            "unit": "шт",
            "quantity": "2",
            "mass_kg": None,
            "note": "ok",
        },
        {
            "sample_id": "sample",
            "page_ref": "1",
            "section": "ПВ7",
            "subsection": "КИПиА",
            "pos": None,
            "name": "Кабель",
            "type_mark": None,
            "code": None,
            "manufacturer": None,
            "unit": "м",
            "quantity": "3,5",
            "mass_kg": None,
            "note": None,
        },
    ]
    path = tmp_path / "rows.jsonl"
    path.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8"
    )
    return path


def test_kp_display_name_appends_type_mark_when_missing() -> None:
    row = SpecificationRow(
        sample_id="sample",
        page_ref="1",
        section=None,
        subsection=None,
        pos="1",
        name="Датчик",
        type_mark="ТС-1",
        code=None,
        manufacturer=None,
        unit="шт",
        quantity=None,
        mass_kg=None,
        note=None,
    )

    assert kp_display_name(row) == "Датчик ТС-1"


def test_linear_quantity_to_surface_area_converts_round_and_rectangular_ducts() -> None:
    assert str(linear_quantity_to_surface_area("Воздуховод Ø100", Decimal("22.0"))) == "6.9"
    assert str(linear_quantity_to_surface_area("Воздуховод 200x150", Decimal("22.0"))) == "15.4"


def test_commercial_unit_quantity_converts_customer_19_formula_items() -> None:
    cases = [
        ("Воздуховод круглого сечения D=200 L=1250", "шт", Decimal("10"), "7.9"),
        ("Воздуховод прямоугольного сечения 200x300 L=400", "шт", Decimal("10"), "4.0"),
        ("Отвод круглого сечения D=1000 угол=90 R=1000 р=0", "шт", Decimal("1"), "4.9"),
        ("Отвод прямоугольного сечения 200x200 угол=90 R=100 р=50", "шт", Decimal("5"), "1.7"),
        ("Переход круглого сечения D=900 d1=1400 L=800 р=100", "шт", Decimal("1000"), "3749.6"),
        (
            "Переход прямоугольного сечения 400x600 200x100 L=900 р=50",
            "шт",
            Decimal("1000"),
            "1331.1",
        ),
        (
            "Переход с прямоугольного на круглое сечение 800x900 D=500 L=800 р=100",
            "шт",
            Decimal("1000"),
            "2540.8",
        ),
        ("Тройник круглого сечения D=800 d1=800 L=1000 l2=1000", "шт", Decimal("2"), "10.0"),
        (
            "Тройник круглого сечения с прямоугольной врезкой D=400 100x300 L=400 l2=100",
            "шт",
            Decimal("1"),
            "0.6",
        ),
        ("Тройник прямоугольного сечения 300x400 100x100 L=200 l2=100", "шт", Decimal("5"), "1.6"),
        (
            "Тройник прямоугольного сечения c круглой врезкой 300x400 D=100 L=200 l2=100",
            "шт",
            Decimal("5"),
            "1.5",
        ),
        ("Заглушка круглого сечения D=280 р=50", "шт", Decimal("5"), "0.5"),
        ("Заглушка прямоугольного сечения 100x100 р=50", "шт", Decimal("5"), "0.2"),
        (
            "Переход (Утка) круглого сечения D=250 L=700 смещение=300 р=50",
            "шт",
            Decimal("5"),
            "3.4",
        ),
        (
            "Переход (Утка) прямоугольного сечения 500x400 L=900 смещение=900 р=50",
            "шт",
            Decimal("1000"),
            "2653.6",
        ),
        (
            "Крестовина круглого сечения D=1000 d2=400 d3=400 L=1200 l2=200 l3=200",
            "шт",
            Decimal("1000"),
            "4069.7",
        ),
        (
            "Крестовина прямоугольного сечения 200x300 100x100 100x100 L=800 l2=400 l3=400",
            "шт",
            Decimal("5"),
            "5.5",
        ),
        ("Врезка прямая круглая D=300 L=100", "шт", Decimal("30"), "2.8"),
        ("Врезка прямая прямоугольная 600x100 L=200", "шт", Decimal("5"), "0.8"),
    ]

    for name, unit, quantity, expected_quantity in cases:
        converted_unit, converted_quantity = commercial_unit_quantity(name, unit, quantity)
        assert converted_unit == "м2"
        assert str(converted_quantity) == expected_quantity


def test_commercial_unit_quantity_accepts_normalized_round_dimension_symbols() -> None:
    cases = [
        ("Отвод круглого сечения Ø1000 угол=90 R=1000 р=0", Decimal("1"), "4.9"),
        ("Заглушка круглого сечения Ø280 р=50", Decimal("5"), "0.5"),
        ("Врезка прямая круглая Ø300 L=100", Decimal("30"), "2.8"),
    ]

    for name, quantity, expected_quantity in cases:
        converted_unit, converted_quantity = commercial_unit_quantity(name, "шт", quantity)
        assert converted_unit == "м2"
        assert str(converted_quantity) == expected_quantity


def test_commercial_unit_quantity_does_not_convert_already_area_or_non_duct_linear_items() -> None:
    assert commercial_unit_quantity(
        "Переход круглого сечения D=900 d1=1400 L=800 р=100",
        "м2",
        Decimal("12.3"),
    ) == ("м2", Decimal("12.3"))
    assert commercial_unit_quantity("Труба стальная Ø100", "п.м", Decimal("22")) == (
        "п.м",
        Decimal("22"),
    )


def test_commercial_unit_quantity_parses_explicit_meter_lengths() -> None:
    converted_unit, converted_quantity = commercial_unit_quantity(
        "Воздуховод круглого сечения D=200 L=1,25м",
        "шт",
        Decimal("10"),
    )

    assert converted_unit == "м2"
    assert str(converted_quantity) == "7.9"


def test_build_kp_output_rows_uses_context_header_for_duct_children() -> None:
    rows = [
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section="ПВ7",
            subsection="КИПиА",
            pos="23",
            name="Воздуховод круглого сечения ГОСТ 14918-80",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit=None,
            quantity=None,
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section="ПВ7",
            subsection="КИПиА",
            pos=None,
            name="ф100",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit="п.м",
            quantity=Decimal("22.0"),
            mass_kg=None,
            note=None,
        ),
    ]

    item_rows = [row for row in build_kp_output_rows(rows) if row.kind == "item"]

    assert len(item_rows) == 1
    assert item_rows[0].pos == "23"
    assert item_rows[0].display_name == "Воздуховод круглого сечения Ø100"
    assert item_rows[0].unit == "м2"
    assert item_rows[0].quantity == Decimal("6.9")


def test_build_kp_output_rows_converts_repaired_duct_dimensions_to_area() -> None:
    rows = [
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section="ПВ7",
            subsection=None,
            pos="23",
            name="Воздуховод из оцинкованной стали",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit=None,
            quantity=None,
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section="ПВ7",
            subsection=None,
            pos=None,
            name="₽100",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit="п.м",
            quantity=Decimal("22.0"),
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section="ПВ7",
            subsection=None,
            pos=None,
            name="0100x050",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit="п.м",
            quantity=Decimal("22.0"),
            mass_kg=None,
            note=None,
        ),
    ]

    item_rows = [row for row in build_kp_output_rows(rows) if row.kind == "item"]

    assert item_rows[0].display_name == "Воздуховод из оцинкованной стали Ø100"
    assert item_rows[0].unit == "м2"
    assert item_rows[0].quantity == Decimal("6.9")
    assert item_rows[1].display_name == "Воздуховод из оцинкованной стали Ø100x50"
    assert item_rows[1].unit == "м2"
    assert item_rows[1].quantity == Decimal("6.6")


def test_build_kp_output_rows_does_not_repeat_composition_header() -> None:
    rows = [
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section="Общеобменная вентиляция",
            subsection="Жилая часть",
            pos="В9.1",
            name="Вытяжная установка, в составе:",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit=None,
            quantity=None,
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section="Общеобменная вентиляция",
            subsection="Жилая часть",
            pos=None,
            name="- вентилятор, 0,07 кВт",
            type_mark="Канал-ВЕНТ-100",
            code=None,
            manufacturer="Веза",
            unit="шт",
            quantity=Decimal("1"),
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section="Общеобменная вентиляция",
            subsection="Жилая часть",
            pos=None,
            name="- шумоглушитель",
            type_mark="Канал-ГКК-100-600",
            code=None,
            manufacturer="Веза",
            unit="шт",
            quantity=Decimal("1"),
            mass_kg=None,
            note=None,
        ),
    ]

    item_rows = [row for row in build_kp_output_rows(rows) if row.kind == "item"]

    assert [row.pos for row in item_rows] == ["В9.1", None, None]
    assert item_rows[0].display_name == "Вытяжная установка, в составе:"
    assert item_rows[1].display_name == "- вентилятор, 0,07 кВт Канал-ВЕНТ-100 Веза"
    assert item_rows[2].display_name == "- шумоглушитель Канал-ГКК-100-600 Веза"
    assert all(
        not (row.display_name or "").startswith("Вытяжная установка, в составе: -")
        for row in item_rows[1:]
    )


def test_build_kp_output_rows_skips_label_only_rows_without_previous_item() -> None:
    rows = [
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section="В20",
            subsection=None,
            pos=None,
            name="В20",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit=None,
            quantity=None,
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section="В20",
            subsection=None,
            pos="1",
            name="Вентилятор",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit="шт",
            quantity=Decimal("1"),
            mass_kg=None,
            note=None,
        ),
    ]

    item_rows = [row for row in build_kp_output_rows(rows) if row.kind == "item"]

    assert len(item_rows) == 1
    assert item_rows[0].display_name == "Вентилятор"
    assert item_rows[0].unit == "шт"


def test_build_kp_output_rows_preserves_named_header_under_group() -> None:
    rows = [
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section=None,
            subsection="Блок 5",
            pos=None,
            name="Теплоснабжение (обвязка горизонтальных завес Guard pro 150W)",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit=None,
            quantity=None,
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section=None,
            subsection=None,
            pos="1",
            name="Завеса воздушная водяная",
            type_mark="GUARD PRO 150W",
            code=None,
            manufacturer="Sonniger",
            unit="wm",
            quantity=Decimal("84"),
            mass_kg=None,
            note=None,
        ),
    ]

    kp_rows = build_kp_output_rows(rows)

    assert [(row.kind, row.label) for row in kp_rows[:2]] == [
        ("group", "Блок 5"),
        ("group", "Теплоснабжение (обвязка горизонтальных завес Guard pro 150W)"),
    ]
    item_rows = [row for row in kp_rows if row.kind == "item"]
    assert item_rows[0].display_name == "Завеса воздушная водяная GUARD PRO 150W Sonniger"
    assert item_rows[0].unit == "шт"


def test_build_kp_output_rows_repairs_ocr_dimension_text() -> None:
    rows = [
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section=None,
            subsection=None,
            pos="10",
            name="Отвод стальной",
            type_mark="ГОСТ 17375-2001",
            code=None,
            manufacturer=None,
            unit=None,
            quantity=None,
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section=None,
            subsection=None,
            pos=None,
            name="025",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit="wm",
            quantity=Decimal("76"),
            mass_kg=None,
            note=None,
        ),
    ]

    item_rows = [row for row in build_kp_output_rows(rows) if row.kind == "item"]

    assert item_rows[0].display_name == "Отвод стальной ГОСТ 17375-2001 Ø25"
    assert item_rows[0].unit == "шт"


def test_build_kp_output_rows_repairs_latex_ocr_symbols() -> None:
    rows = [
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section=None,
            subsection=None,
            pos="4",
            name="Кран шаровой латунный $ \\Phi $15",
            type_mark="BVR-R- $ \\Phi $15",
            code=None,
            manufacturer="Danfoss",
            unit="$ \\omega $m",
            quantity=Decimal("4"),
            mass_kg=None,
            note=None,
        )
    ]

    item_rows = [row for row in build_kp_output_rows(rows) if row.kind == "item"]

    assert item_rows[0].display_name == "Кран шаровой латунный Ø15 BVR-R-Ø15 Danfoss"
    assert item_rows[0].unit == "шт"


def test_build_kp_output_rows_marks_multiline_crossed_out_item_red() -> None:
    rows = [
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section=None,
            subsection=None,
            pos="6",
            name="Смесительный узел",
            type_mark="SUS20-4,0 TS",
            code=None,
            manufacturer=None,
            unit="wm",
            quantity=Decimal("38"),
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section=None,
            subsection=None,
            pos=None,
            name="гибкие подводки $ \\Phi $15",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit=None,
            quantity=None,
            mass_kg=None,
            note=None,
            status="crossed_out",
        ),
    ]

    item_rows = [row for row in build_kp_output_rows(rows) if row.kind == "item"]

    assert item_rows[0].display_name == "Смесительный узел SUS20-4,0 TS гибкие подводки Ø15"
    assert item_rows[0].status == "crossed_out"


def test_build_kp_output_rows_does_not_glue_new_position_heading_to_previous_item() -> None:
    rows = [
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section=None,
            subsection=None,
            pos="1.3",
            name="Воздуховод",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit="м/м²",
            quantity=None,
            quantity_raw="1408/70,7",
            mass_kg=None,
            note="Воздушные затворы",
        ),
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section=None,
            subsection=None,
            pos="1.4",
            name="Отводы прямоугольного сечения",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit=None,
            quantity=None,
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section=None,
            subsection=None,
            pos=None,
            name="160 (δ=1 мм)",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit="шт./м²",
            quantity=Decimal("455"),
            mass_kg=Decimal("56.14"),
            note=None,
        ),
    ]

    kp_rows = build_kp_output_rows(rows)
    item_rows = [row for row in kp_rows if row.kind == "item"]

    assert item_rows[0].display_name == "Воздуховод"
    assert item_rows[0].unit == "м2"
    assert item_rows[0].quantity == Decimal("70.7")
    assert all("1.4 Отводы" not in (row.display_name or "") for row in item_rows)
    assert any(
        row.kind == "group" and row.label == "1.4 Отводы прямоугольного сечения" for row in kp_rows
    )


def test_build_kp_output_rows_defers_prefix_header_to_next_quantity_row() -> None:
    rows = [
        SpecificationRow(
            sample_id="sample",
            page_ref="2",
            section="Отопление",
            subsection=None,
            pos=None,
            name="Монтажный набор для регулирования",
            type_mark="Unibox RTL",
            code=None,
            manufacturer="Oventrop",
            unit="шт.",
            quantity=Decimal("1"),
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="sample",
            page_ref="2",
            section="Отопление",
            subsection=None,
            pos=None,
            name="Плита шумо- теплоизоляционная Energofloor под водяной",
            type_mark="TACKER AL 20/1-4,8 DES-",
            code="EFRP2014/8TKR",
            manufacturer=None,
            unit=None,
            quantity=None,
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="sample",
            page_ref="2",
            section="Отопление",
            subsection=None,
            pos=None,
            name="теплый пол, с алюминизированным полимерным покрытием",
            type_mark="sg",
            code="AL",
            manufacturer="Rols-Isomarket",
            unit="м²",
            quantity=Decimal("33"),
            mass_kg=None,
            note=None,
        ),
    ]

    item_rows = [row for row in build_kp_output_rows(rows) if row.kind == "item"]

    assert item_rows[0].display_name == "Монтажный набор для регулирования Unibox RTL Oventrop"
    assert item_rows[0].quantity == Decimal("1")
    assert item_rows[1].display_name == (
        "Плита шумо- теплоизоляционная Energofloor под водяной "
        "TACKER AL 20/1-4,8 DES- EFRP2014/8TKR "
        "теплый пол, с алюминизированным полимерным покрытием sg AL Rols-Isomarket"
    )
    assert item_rows[1].quantity == Decimal("33")


def test_build_kp_output_rows_keeps_short_header_as_group_not_previous_item_text() -> None:
    rows = [
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section=None,
            subsection=None,
            pos="3",
            name="Компрессорно-конденсаторный блок SKB-14WC/SF",
            type_mark=None,
            code=None,
            manufacturer="СВОК",
            unit="шт",
            quantity=Decimal("1"),
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="sample",
            page_ref="1",
            section=None,
            subsection=None,
            pos=None,
            name="КИПиА",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit=None,
            quantity=None,
            mass_kg=None,
            note=None,
        ),
    ]

    kp_rows = build_kp_output_rows(rows)

    assert kp_rows[0].kind == "item"
    assert kp_rows[0].display_name == "Компрессорно-конденсаторный блок SKB-14WC/SF СВОК"
    assert kp_rows[1].kind == "group"
    assert kp_rows[1].label == "КИПиА"


def test_build_kp_output_rows_uses_literal_formula_quantity_for_reserve_expression() -> None:
    row = SpecificationRow(
        sample_id="sample",
        page_ref="1",
        section=None,
        subsection=None,
        pos="1",
        name="Канальный вентилятор",
        type_mark=None,
        code=None,
        manufacturer=None,
        unit="шт",
        quantity=None,
        quantity_raw="18+(18-резерв)",
        mass_kg=None,
        note=None,
    )

    item_rows = [row for row in build_kp_output_rows([row]) if row.kind == "item"]

    assert item_rows[0].quantity == Decimal("36")
    assert item_rows[0].calc_quantity == Decimal("36")


def test_export_jsonl_to_kp_xlsx_uses_area_part_of_compound_quantity_for_formula(
    tmp_path: Path,
) -> None:
    template = _write_kp_template(tmp_path)
    input_jsonl = tmp_path / "rows.jsonl"
    input_jsonl.write_text(
        json.dumps(
            {
                "sample_id": "sample",
                "page_ref": "1",
                "section": None,
                "subsection": None,
                "pos": "1.1",
                "name": "Воздуховод",
                "type_mark": None,
                "code": None,
                "manufacturer": None,
                "unit": "м/м²",
                "quantity": "55/17,27",
                "mass_kg": None,
                "note": None,
            },
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    output = tmp_path / "out.xlsx"

    export_jsonl_to_kp_xlsx(
        input_jsonl_path=input_jsonl, output_path=output, template_path=template
    )

    worksheet = load_workbook(output, data_only=False)[MAIN_SHEET_NAME]
    assert worksheet["C11"].value == "м2"
    assert worksheet["D11"].value == 17.27
    assert worksheet["G11"].value == "=E11*D11"
    assert worksheet["H11"].value == "=F11*D11"


def test_export_jsonl_to_kp_xlsx_expands_long_rows_for_wrapped_text(tmp_path: Path) -> None:
    template = _write_kp_template(tmp_path)
    input_jsonl = tmp_path / "rows.jsonl"
    input_jsonl.write_text(
        json.dumps(
            {
                "sample_id": "sample",
                "page_ref": "1",
                "section": None,
                "subsection": None,
                "pos": "1",
                "name": (
                    "Вентилятор радиальный дымоудаления производительностью L=22740м³/ч, "
                    "Pс=350Па, с электродвигателем N=11кВт, n=1500об/мин. комплектно "
                    "с щитом управления и автоматики"
                ),
                "type_mark": "VR 80-75-710-D-2h/400-11/1500-ПР0-У1",
                "code": None,
                "manufacturer": "LUFTKON",
                "unit": "шт",
                "quantity": "1",
                "mass_kg": None,
                "note": "КП №DEMO-0042 тел. DEMO-PHONE-0000 от 01.02.2030",
            },
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    output = tmp_path / "out.xlsx"

    export_jsonl_to_kp_xlsx(
        input_jsonl_path=input_jsonl, output_path=output, template_path=template
    )

    worksheet = load_workbook(output, data_only=False)[MAIN_SHEET_NAME]
    assert worksheet["B11"].alignment.wrap_text is True
    assert worksheet["I11"].alignment.wrap_text is True
    assert worksheet.row_dimensions[11].height is not None
    assert worksheet.row_dimensions[11].height >= 45


def test_export_jsonl_to_kp_xlsx_uses_template_column_width_for_row_height(
    tmp_path: Path,
) -> None:
    template = _write_kp_template(tmp_path)
    workbook = load_workbook(template)
    worksheet = workbook[MAIN_SHEET_NAME]
    worksheet.column_dimensions["B"].width = 120
    workbook.save(template)
    input_jsonl = tmp_path / "rows.jsonl"
    input_jsonl.write_text(
        json.dumps(
            {
                "sample_id": "sample",
                "page_ref": "1",
                "section": None,
                "subsection": None,
                "pos": "1",
                "name": "Канальный вентилятор N=2,2квт n=1280об/мин 380/50 с регулятором скорости",
                "type_mark": None,
                "code": None,
                "manufacturer": None,
                "unit": "шт",
                "quantity": "1",
                "mass_kg": None,
                "note": None,
            },
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    output = tmp_path / "out.xlsx"

    export_jsonl_to_kp_xlsx(
        input_jsonl_path=input_jsonl, output_path=output, template_path=template
    )

    worksheet = load_workbook(output, data_only=False)[MAIN_SHEET_NAME]
    assert worksheet.row_dimensions[11].height == 15


def test_export_jsonl_to_kp_xlsx_preserves_template_and_writes_body(tmp_path: Path) -> None:
    template = _write_kp_template(tmp_path)
    _write_kp_header_template(tmp_path)
    input_jsonl = _write_rows_jsonl(tmp_path)
    output = tmp_path / "out.xlsx"

    written_path, row_count = export_jsonl_to_kp_xlsx(
        input_jsonl_path=input_jsonl,
        output_path=output,
        template_path=template,
    )

    assert written_path == output
    assert row_count == 2
    workbook = load_workbook(output, data_only=False)
    worksheet = workbook[MAIN_SHEET_NAME]
    assert workbook.sheetnames == [MAIN_SHEET_NAME]
    assert worksheet["A2"].value == "согласовано                          м.п."
    assert worksheet["H2"].value == "утверждаю                        м.п."
    assert worksheet["A3"].value == "Коммерческое предложение по спецификации"
    assert worksheet["A4"].value is None
    assert "A3:I3" in {str(merged_range) for merged_range in worksheet.merged_cells.ranges}
    assert worksheet["B7"].value == "Наименование"
    assert worksheet["B11"].value == "ПВ7"
    assert worksheet["B12"].value == "КИПиА"
    assert worksheet["A13"].value == "4"
    assert worksheet["B13"].value == "Датчик ТС-1 СВОК"
    assert worksheet["C13"].value == "шт"
    assert worksheet["D14"].value == 3.5
    assert worksheet["D13"].number_format == "0.00"
    assert worksheet["D14"].number_format == "0.00"
    assert worksheet["I13"].value == "ok"
    assert worksheet["G13"].value == "=E13*D13"
    assert worksheet["H13"].value == "=F13*D13"
    assert all(worksheet.cell(15, column).value is None for column in range(1, 10))
    assert worksheet["C16"].value == "Итого стоимость оборудования и материалов:"
    assert worksheet["H16"].value == "=SUM(G11:G14)"
    assert worksheet["A16"].border.top.style == "thin"
    assert worksheet["C16"].border.top.style == "thin"
    assert worksheet["H16"].border.top.style == "thin"
    assert worksheet["H19"].value == "=SUM(H11:H14)"
    assert worksheet["H20"].value == 100000
    assert worksheet["H21"].value == 1200000
    assert worksheet.max_column == 9
    assert worksheet.max_row == 23
    assert worksheet.print_area == "'Quote'!$A$1:$I$23"
    assert "C16:G16" in {str(merged_range) for merged_range in worksheet.merged_cells.ranges}
    assert "K" not in worksheet.column_dimensions


def test_export_jsonl_to_kp_xlsx_marks_crossed_out_rows_red(tmp_path: Path) -> None:
    template = _write_kp_template(tmp_path)
    input_jsonl = tmp_path / "rows.jsonl"
    input_jsonl.write_text(
        json.dumps(
            {
                "sample_id": "sample",
                "page_ref": "1",
                "section": None,
                "subsection": None,
                "pos": "6",
                "name": "Исключенная строка",
                "type_mark": None,
                "code": None,
                "manufacturer": None,
                "unit": "шт",
                "quantity": "1",
                "mass_kg": None,
                "note": None,
                "status": "crossed_out",
            },
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    output = tmp_path / "out.xlsx"

    export_jsonl_to_kp_xlsx(
        input_jsonl_path=input_jsonl, output_path=output, template_path=template
    )

    worksheet = load_workbook(output)[MAIN_SHEET_NAME]
    assert worksheet["B11"].value == "Исключенная строка"
    assert worksheet["B11"].font.color.rgb == "00FF0000"


def test_export_jsonl_to_kp_xlsx_escapes_untrusted_formula_text(tmp_path: Path) -> None:
    template = _write_kp_template(tmp_path)
    input_jsonl = tmp_path / "rows.jsonl"
    input_jsonl.write_text(
        json.dumps(
            {
                "sample_id": "untrusted",
                "page_ref": "1",
                "section": None,
                "subsection": None,
                "pos": "@position",
                "name": '=HYPERLINK("https://example.invalid","click")',
                "type_mark": None,
                "code": None,
                "manufacturer": None,
                "unit": "+cmd",
                "quantity": "-command",
                "mass_kg": None,
                "note": "-unsafe",
            }
        )
        + "\n",
        encoding="utf-8",
    )
    output = tmp_path / "safe.xlsx"

    export_jsonl_to_kp_xlsx(
        input_jsonl_path=input_jsonl, output_path=output, template_path=template
    )

    worksheet = load_workbook(output, data_only=False)[MAIN_SHEET_NAME]
    for coordinate in ("A11", "B11", "C11", "D11", "I11"):
        assert worksheet[coordinate].data_type != "f"
    assert worksheet["A11"].value == "'@position"
    assert worksheet["B11"].value.startswith("'=")
    assert worksheet["C11"].value == "'+cmd"
    assert worksheet["D11"].value == "'-command"
    assert worksheet["I11"].value == "'-unsafe"
    assert worksheet["G11"].value is None
    assert worksheet["H11"].value is None
