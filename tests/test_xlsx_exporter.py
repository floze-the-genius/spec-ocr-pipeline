from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from spec_pipeline.exporters.xlsx import (
    SHEET_NAME,
    default_xlsx_output_path,
    export_jsonl_to_xlsx,
    export_sample_to_xlsx,
    write_rows_to_xlsx,
)
from spec_pipeline.models import SpecificationRow
from spec_pipeline.samples import load_sample_manifest


def _write_manifest(tmp_path: Path) -> Path:
    tsv_path = tmp_path / "sample.tsv"
    tsv_path.write_text(
        "page\tsection\tsubsection\tpos\tname\ttype_mark\tcode\tmanufacturer\tunit\tquantity\tmass_kg\tnote\n"
        "001\tПВ7\t\t\t\t\t\t\t\t\t\t\n"
        "001\tПВ7\tКИПиА\t1\tДатчик\tТС-1\t\tСВОК\tшт\t12,5\t1,000\tok\n",
        encoding="utf-8",
    )
    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(
        json.dumps(
            {
                "manifest_version": "1",
                "name": "xlsx-test",
                "samples": [
                    {
                        "sample_id": "sample_1",
                        "title": "Sample 1",
                        "source_pdf": str(tmp_path / "input.pdf"),
                        "ground_truth_tsv": str(tsv_path),
                        "expected_xlsx": str(tmp_path / "expected.xlsx"),
                        "visual_xlsx": str(tmp_path / "visual.xlsx"),
                        "page_field": "page",
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return manifest_path


def test_write_rows_to_xlsx_creates_single_sheet_with_header(tmp_path: Path) -> None:
    output_path = tmp_path / "out.xlsx"
    rows = [
        SpecificationRow(
            sample_id="spec_1",
            page_ref="001",
            section="ПВ7",
            subsection=None,
            pos="1",
            name="Установка",
            type_mark=None,
            code=None,
            manufacturer="СВОК",
            unit="шт",
            quantity=Decimal("1"),
            mass_kg=None,
            note=None,
        )
    ]

    write_rows_to_xlsx(rows, output_path)

    workbook = load_workbook(output_path)
    worksheet = workbook[SHEET_NAME]
    assert workbook.sheetnames == [SHEET_NAME]
    assert worksheet[1][0].value == "№ п/п"
    assert worksheet[1][1].value == "Наименование"
    assert worksheet[2][1].value == "ПВ7"
    assert worksheet[3][1].value == "Установка СВОК"


def test_write_rows_to_xlsx_includes_grouping_rows(tmp_path: Path) -> None:
    output_path = tmp_path / "out.xlsx"
    rows = [
        SpecificationRow(
            sample_id="spec_1",
            page_ref="001",
            section="ПВ7",
            subsection=None,
            pos=None,
            name=None,
            type_mark=None,
            code=None,
            manufacturer=None,
            unit=None,
            quantity=None,
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="spec_1",
            page_ref="001",
            section="ПВ7",
            subsection="КИПиА",
            pos="1",
            name="Датчик",
            type_mark="ТС-1",
            code=None,
            manufacturer="СВОК",
            unit="шт",
            quantity=Decimal("12.5"),
            mass_kg=None,
            note=None,
        ),
    ]

    write_rows_to_xlsx(rows, output_path)

    worksheet = load_workbook(output_path)[SHEET_NAME]
    assert worksheet.max_row == 4
    assert worksheet[2][1].value == "ПВ7"
    assert worksheet[3][1].value == "КИПиА"
    assert worksheet[4][1].value == "Датчик ТС-1 СВОК"


def test_write_rows_to_xlsx_preserves_decimal_values(tmp_path: Path) -> None:
    output_path = tmp_path / "out.xlsx"
    rows = [
        SpecificationRow(
            sample_id="spec_1",
            page_ref="001",
            section="ПВ7",
            subsection="КИПиА",
            pos="1",
            name="Датчик",
            type_mark="ТС-1",
            code=None,
            manufacturer="СВОК",
            unit="шт",
            quantity=Decimal("12.5"),
            mass_kg=Decimal("1.000"),
            note=None,
        )
    ]

    write_rows_to_xlsx(rows, output_path)

    worksheet = load_workbook(output_path)[SHEET_NAME]
    assert worksheet[4][3].value == 12.5
    assert worksheet[4][3].number_format == "0.00"
    assert worksheet.max_column == 5


def test_write_rows_to_xlsx_preserves_compound_quantity_text(tmp_path: Path) -> None:
    output_path = tmp_path / "out.xlsx"
    rows = [
        SpecificationRow(
            sample_id="spec_1",
            page_ref="001",
            section=None,
            subsection=None,
            pos="1.1",
            name="Воздуховод",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit="м/м²",
            quantity=None,
            quantity_raw="55/17,27",
            mass_kg=None,
            note=None,
        )
    ]

    write_rows_to_xlsx(rows, output_path)

    worksheet = load_workbook(output_path)[SHEET_NAME]
    assert worksheet[2][3].value == "55/17,27"


def test_write_rows_to_xlsx_marks_crossed_out_rows_red(tmp_path: Path) -> None:
    output_path = tmp_path / "out.xlsx"
    rows = [
        SpecificationRow(
            sample_id="spec_1",
            page_ref="001",
            section=None,
            subsection=None,
            pos="1",
            name="Исключенная строка",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit="шт",
            quantity=Decimal("1"),
            mass_kg=None,
            note=None,
            status="crossed_out",
        )
    ]

    write_rows_to_xlsx(rows, output_path)

    worksheet = load_workbook(output_path)[SHEET_NAME]
    assert worksheet[2][1].value == "Исключенная строка"
    assert worksheet[2][1].font.color.rgb == "00FF0000"


def test_write_rows_to_xlsx_repairs_ocr_dimension_and_unit_text(tmp_path: Path) -> None:
    output_path = tmp_path / "out.xlsx"
    rows = [
        SpecificationRow(
            sample_id="spec_1",
            page_ref="001",
            section=None,
            subsection=None,
            pos=None,
            name="₽50₽25",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit="w/m",
            quantity=Decimal("168"),
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="spec_1",
            page_ref="001",
            section=None,
            subsection=None,
            pos=None,
            name="057x3.5",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit="n.m.",
            quantity=Decimal("96"),
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="spec_1",
            page_ref="001",
            section=None,
            subsection=None,
            pos=None,
            name="Кран $ \\Phi $15",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit="$ \\omega $m",
            quantity=Decimal("4"),
            mass_kg=None,
            note=None,
        ),
    ]

    write_rows_to_xlsx(rows, output_path)

    worksheet = load_workbook(output_path)[SHEET_NAME]
    assert worksheet[2][1].value == "Ø50x25"
    assert worksheet[2][2].value == "шт"
    assert worksheet[3][1].value == "Ø57x3.5"
    assert worksheet[3][2].value == "п.м."
    assert worksheet[4][1].value == "Кран Ø15"
    assert worksheet[4][2].value == "шт"


def test_write_rows_to_xlsx_marks_multiline_crossed_out_item_red(tmp_path: Path) -> None:
    output_path = tmp_path / "out.xlsx"
    rows = [
        SpecificationRow(
            sample_id="spec_1",
            page_ref="001",
            section=None,
            subsection=None,
            pos="6",
            name="Смесительный узел",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit="шт",
            quantity=Decimal("38"),
            mass_kg=None,
            note=None,
        ),
        SpecificationRow(
            sample_id="spec_1",
            page_ref="001",
            section=None,
            subsection=None,
            pos=None,
            name="гибкие подводки $ \\Phi $15",
            type_mark=None,
            code=None,
            manufacturer=None,
            unit=None,
            quantity=None,
            mass_kg=None,
            note=None,
            status="crossed_out",
        ),
    ]

    write_rows_to_xlsx(rows, output_path)

    worksheet = load_workbook(output_path)[SHEET_NAME]
    assert worksheet[2][1].font.color.rgb == "00FF0000"
    assert worksheet[3][1].value == "гибкие подводки Ø15"
    assert worksheet[3][1].font.color.rgb == "00FF0000"


def test_export_sample_to_xlsx_uses_default_artifact_path(tmp_path: Path) -> None:
    manifest_path = _write_manifest(tmp_path)
    manifest = load_sample_manifest(manifest_path)

    output_path, row_count = export_sample_to_xlsx(
        sample=manifest.samples[0],
        project_root=tmp_path,
    )

    assert row_count == 2
    assert output_path == default_xlsx_output_path(tmp_path, "sample_1")
    assert output_path.exists()
    workbook = load_workbook(output_path)
    assert workbook.sheetnames == [SHEET_NAME]


def test_export_jsonl_to_xlsx_writes_rows_without_manifest(tmp_path: Path) -> None:
    input_jsonl = tmp_path / "rows.jsonl"
    input_jsonl.write_text(
        json.dumps(
            {
                "sample_id": "ai_sample",
                "page_ref": "1",
                "section": None,
                "subsection": None,
                "pos": "1",
                "name": "Котел",
                "type_mark": "ALIXIA 24FF",
                "code": None,
                "manufacturer": "DemoHeat",
                "unit": "компл.",
                "quantity": "1",
                "mass_kg": None,
                "note": None,
            },
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    output_path = tmp_path / "rows.xlsx"

    written_path, row_count = export_jsonl_to_xlsx(
        input_jsonl_path=input_jsonl, output_path=output_path
    )

    assert written_path == output_path
    assert row_count == 1
    worksheet = load_workbook(output_path)[SHEET_NAME]
    assert worksheet[2][1].value == "Котел ALIXIA 24FF DemoHeat"


def test_write_rows_to_xlsx_escapes_untrusted_formula_text(tmp_path: Path) -> None:
    output_path = tmp_path / "safe.xlsx"
    rows = [
        SpecificationRow(
            sample_id="untrusted",
            page_ref="1",
            section=None,
            subsection=None,
            pos="@position",
            name='=HYPERLINK("https://example.invalid","click")',
            type_mark=None,
            code=None,
            manufacturer=None,
            unit="+cmd",
            quantity=None,
            mass_kg=None,
            note="-unsafe",
        )
    ]

    write_rows_to_xlsx(rows, output_path)

    worksheet = load_workbook(output_path, data_only=False)[SHEET_NAME]
    for cell in worksheet[2]:
        if cell.value is not None:
            assert cell.data_type != "f"
    assert worksheet["A2"].value == "'@position"
    assert worksheet["B2"].value.startswith("'=")
    assert worksheet["C2"].value == "'+cmd"
    assert worksheet["E2"].value == "'-unsafe"
